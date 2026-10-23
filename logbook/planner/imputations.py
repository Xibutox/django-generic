"""The imputation export: time logged, spread over the distribution
matrix (``ImputationShare``) and packed into days of 7.4 hours, as an
Excel file - the old planner's "Exporter la repartition des imputations
en Excel", same rules, same sheets.

1. The time selected, counted on days of 8 hours, is converted to days
   of 7.4 hours (to the second, rounded half up).
2. Each line of the matrix gets its percentage of it, rounded down to
   the second; the seconds left go one by one to the largest
   remainders.
3. The shares fill days of 7.4 hours, one after the other: a share
   running past the end of a day goes on the next.
"""

from __future__ import annotations

from decimal import ROUND_FLOOR, ROUND_HALF_UP, Decimal
from io import BytesIO
from typing import Any, Iterable

from django.utils import timezone

SOURCE_HOURS_PER_DAY = Decimal("8")
OUTPUT_HOURS_PER_DAY = Decimal("7.4")
SECONDS_PER_DAY = int(OUTPUT_HOURS_PER_DAY * Decimal("3600"))


class ExportError(Exception):
    """What stops an export, said to the reader."""


def convert(seconds: int) -> int:
    """8-hour-day time on a 7.4-hour-day basis."""
    converted = Decimal(seconds) * OUTPUT_HOURS_PER_DAY / SOURCE_HOURS_PER_DAY

    return int(converted.quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def check_matrix(shares: list[Any]) -> list[Decimal]:
    if not shares:
        raise ExportError("Aucune matrice d'imputation n'est configuree.")

    percentages = [Decimal(str(share.percentage)) for share in shares]

    if any(percentage < 0 for percentage in percentages):
        raise ExportError("Les pourcentages ne peuvent pas etre negatifs.")

    total = sum(percentages, start=Decimal("0"))

    if abs(total - Decimal("100")) > Decimal("0.000001"):
        raise ExportError(
            "La somme des pourcentages doit etre egale a 100 %. "
            f"Valeur actuelle : {total} %."
        )

    return percentages


def allocate(
    total_seconds: int, numbers: list[str], percentages: list[Decimal]
) -> list[dict[str, Any]]:
    allocations = []
    allocated = 0

    for number, percentage in zip(numbers, percentages):
        exact = Decimal(total_seconds) * percentage / Decimal("100")
        seconds = int(exact.to_integral_value(rounding=ROUND_FLOOR))
        allocations.append(
            {
                "numero": number,
                "seconds": seconds,
                "remainder": exact - Decimal(seconds),
            }
        )
        allocated += seconds

    missing = total_seconds - allocated
    ranked = sorted(
        allocations, key=lambda item: item["remainder"], reverse=True
    )

    for index in range(missing):
        ranked[index % len(ranked)]["seconds"] += 1

    return [
        {"numero": item["numero"], "seconds": item["seconds"]}
        for item in allocations
        if item["seconds"] > 0
    ]


def into_days(allocations: list[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = []
    day = 1
    used = 0

    for allocation in allocations:
        remaining = allocation["seconds"]

        while remaining > 0:
            taken = min(remaining, SECONDS_PER_DAY - used)
            rows.append(
                {"day": day, "seconds": taken, "numero": allocation["numero"]}
            )
            remaining -= taken
            used += taken

            if used == SECONDS_PER_DAY:
                day += 1
                used = 0

    return rows


def safe_text(value: Any) -> str:
    """No formula in a cell: a text starting with = + - @ is quoted."""
    value = str(value)

    return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value


def workbook(entries: Iterable[Any], shares: list[Any]) -> tuple[str, bytes]:
    """The file's name and content for these time entries."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    entries = list(entries)
    percentages = check_matrix(shares)
    # To the second, from the start and the end; a running timer has
    # no time yet.
    source_seconds = int(
        sum(
            (entry.ended_at - entry.started_at).total_seconds()
            for entry in entries
            if entry.ended_at is not None
        )
    )

    if source_seconds <= 0:
        raise ExportError("Le temps total de la selection est nul.")

    total = convert(source_seconds)
    rows = into_days(
        allocate(total, [share.number for share in shares], percentages)
    )
    dates = {timezone.localdate(entry.started_at) for entry in entries}

    book = Workbook()
    sheet = book.active
    sheet.title = "Repartition"
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.append(("Jour", "Heures passees", "Numero"))
    fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)

    for cell in sheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        sheet.append(
            [row["day"], row["seconds"] / 3600, safe_text(row["numero"])]
        )

    for cell in sheet["B"][1:]:
        cell.number_format = '0.00 "h"'

    sheet.auto_filter.ref = sheet.dimensions

    for index, width in {1: 12, 2: 20, 3: 30}.items():
        sheet.column_dimensions[get_column_letter(index)].width = width

    summary = book.create_sheet("Resume")
    summary.sheet_view.showGridLines = False
    summary.append(["Information", "Valeur"])
    summary.append(["Temps selectionne sur base 8 h", source_seconds / 3600])
    summary.append(["Temps converti sur base 7.4 h", total / 3600])
    summary.append(["Nombre de dates selectionnees", len(dates)])
    summary.append(
        ["Nombre de jours utilises", max((r["day"] for r in rows), default=0)]
    )
    summary.append(["Capacite journaliere", float(OUTPUT_HOURS_PER_DAY)])

    for cell in summary[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center")

    for coordinate in ("B2", "B3", "B6"):
        summary[coordinate].number_format = '0.00 "h"'

    summary.column_dimensions["A"].width = 38
    summary.column_dimensions["B"].width = 20

    output = BytesIO()
    book.save(output)
    stamp = timezone.localtime().strftime("%Y%m%d_%H%M%S")

    return f"repartition_imputations_{stamp}.xlsx", output.getvalue()
