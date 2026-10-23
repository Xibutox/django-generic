"""The logbook example must keep working.

``logbook/`` is a project of its own - its settings, its ``manage.py``
- so, like the document manager, it is run as one, in a process of its
own, from its folder, on a test database.
"""

from __future__ import annotations

import os
import subprocess
import sys
import textwrap
from pathlib import Path

LOGBOOK = Path(__file__).resolve().parent.parent / "logbook"

#: Seeded twice, then walked by each kind of person: the note editor,
#: the timer, the timesheet's allocation, the verification.
SCENARIO = textwrap.dedent("""
    import json
    import os

    os.environ["DJANGO_SETTINGS_MODULE"] = "logsite.settings"

    import django
    from django.conf import settings

    django.setup()
    settings.ALLOWED_HOSTS = ["testserver"]

    from django.contrib.auth import get_user_model
    from django.core.management import call_command
    from django.db import connection
    from django.test import Client
    from django.test.utils import setup_test_environment

    setup_test_environment()
    connection.creation.create_test_db(verbosity=0)
    call_command("seed_planner", stdout=open(os.devnull, "w"))
    call_command("seed_planner", stdout=open(os.devnull, "w"))

    from planner.models import Note, NoteLine, Task, TimeEntry

    def signed_in(username):
        client = Client()
        client.force_login(get_user_model().objects.get(username=username))
        return client

    def send(client, method, url, body):
        return getattr(client, method)(
            url, json.dumps(body), content_type="application/json"
        )

    # Seeded once, whatever the number of runs.
    assert Note.objects.count() == 5, Note.objects.count()
    assert Task.objects.count() == 18, Task.objects.count()
    assert TimeEntry.objects.running().count() == 1
    references = set(Task.objects.values_list("reference", flat=True))
    assert {"FNC-0001", "DM-0001", "AUT-0001"} <= references, references

    # Only tickets are critical; a task is about several projects, or
    # none.
    assert not Task.objects.filter(kind="FNC", criticality="").exists()
    assert not Task.objects.exclude(kind="FNC").exclude(
        criticality=""
    ).exists()
    assert Task.objects.get(title="Admin dashboard").projects.count() == 2
    assert not Task.objects.get(
        title="Help the support desk with the VPN"
    ).projects.exists()

    admin, bob, viewer = signed_in("admin"), signed_in("bob"), signed_in("viewer")

    # Every page opens.
    note = Note.objects.get(title="Sprint 42 planning")
    task = Task.objects.get(title="One-page checkout")
    entry = TimeEntry.objects.exclude(ended_at=None).first()
    for url in (
        "/",
        "/planner/note/",
        f"/planner/note/{note.pk}/",
        "/planner/note/add/",
        "/planner/task/",
        f"/planner/task/{task.pk}/",
        "/planner/task/board/",
        "/planner/task/people/",
        "/planner/task/add/",
        f"/planner/project/{Task.objects.get(title='Admin dashboard').projects.first().pk}/",
        "/api/planner/task/charts/by_kind/",
        "/api/planner/timeentry/charts/by_project/",
        "/planner/task/due/",
        "/planner/timeentry/",
        "/planner/timeentry/timesheet/",
        f"/planner/timeentry/{entry.pk}/",
        "/planner/timeentry/add/",
        "/planner/project/",
        "/planner/label/",
    ):
        status = bob.get(url).status_code
        assert status == 200, (url, status)

    # The editor reads the note: its lines, their tasks.
    DOCUMENT = f"/planner/note/{note.pk}/document/"
    read = bob.get(DOCUMENT).json()
    assert read["canEdit"] and read["canAddTasks"]
    kinds = [line["kind"] for line in read["lines"]]
    assert kinds.count("task") == 4, kinds
    first_task = next(line for line in read["lines"] if line["task"])
    assert first_task["task"]["reference"].startswith("FNC-")

    # A line typed as a task makes a task, numbered by its kind; a task
    # line renamed renames its task.
    lines = [
        {k: line[k] for k in ("id", "kind", "indent", "text", "done")}
        for line in read["lines"]
    ]
    lines[4]["text"] = "Login fails after a password reset (Safari)"
    lines.append(
        {"kind": "task", "text": "Export the timesheet", "taskKind": "DM"}
    )
    lines.append(
        {"kind": "task", "text": "Demo to the customer", "taskKind": "AUT"}
    )
    lines.append({"kind": "text", "text": "Just a thought."})
    written = send(bob, "post", DOCUMENT, {"title": note.title, "lines": lines})
    assert written.status_code == 200, written.content
    made = Task.objects.get(title="Export the timesheet")
    assert made.kind == "DM" and made.note == note
    assert made.reference.startswith("DM-")
    assert [p.name for p in made.projects.all()] == ["Website"]
    meeting = Task.objects.get(title="Demo to the customer")
    assert meeting.reference.startswith("AUT-") and not meeting.criticality
    lines_read = written.json()["lines"]
    assert {c["value"] for c in written.json()["criticalities"]} == {
        "minor", "major", "critical"
    }
    fnc = next(
        line["task"] for line in lines_read
        if line["task"] and line["task"]["kind"] == "FNC"
    )
    assert fnc["criticality"] in ("minor", "major", "critical")

    # A ticket's criticality, changed from the note, as the editor does.
    changed = send(
        bob, "patch", fnc["api"], {"criticality": "critical"}
    )
    assert changed.status_code == 200, changed.content
    assert Task.objects.get(pk=fnc["id"]).criticality == "critical"

    # The form refuses a ticket without one.
    from planner.models import Task as T
    from django.core.exceptions import ValidationError
    try:
        T(kind="FNC", title="x").full_clean(exclude=["reference"])
    except ValidationError as error:
        assert "criticality" in error.message_dict
    else:
        raise AssertionError("a ticket without criticality")
    assert Task.objects.filter(
        title="Login fails after a password reset (Safari)"
    ).exists()

    # Gone from the note, untouched: the task goes too.
    lines = [
        {k: line[k] for k in ("id", "kind", "indent", "text", "done")}
        for line in written.json()["lines"]
        if not (line["task"] and line["task"]["id"] == made.pk)
    ]
    assert send(
        bob, "post", DOCUMENT, {"title": note.title, "lines": lines}
    ).status_code == 200
    assert not Task.objects.filter(pk=made.pk).exists()

    # A reader reads, and does not write.
    assert viewer.get(DOCUMENT).json()["canEdit"] is False
    assert send(
        viewer, "post", DOCUMENT, {"title": "Mine", "lines": []}
    ).status_code == 403
    assert send(bob, "post", DOCUMENT, {"lines": "nope"}).status_code == 400

    # The timer: one running at a time, the task moving along.
    todo = Task.objects.get(title="Address autocomplete")
    assert todo.status == "todo"
    started = bob.post(f"/planner/task/{todo.pk}/timer/")
    assert started.status_code == 200, started.content
    assert started.json()["running"] is True
    todo.refresh_from_db()
    assert todo.status == "doing"
    user_bob = get_user_model().objects.get(username="bob")
    assert TimeEntry.objects.running().filter(user=user_bob).count() == 1
    stopped = bob.post(f"/planner/task/{todo.pk}/timer/").json()
    assert stopped["running"] is False
    assert not TimeEntry.objects.running().filter(user=user_bob).exists()
    assert viewer.post(f"/planner/task/{todo.pk}/timer/").status_code == 403

    # The timesheet: the tasks under way are offered, one dropped on a
    # time logs an hour of it, for whoever dropped it.
    SHEET = "/api/planner/timeentry/calendars/timesheet/"
    offered = bob.get(SHEET + "sources/").json()["items"]
    assert offered and all(
        Task.objects.get(pk=item["id"]).status in ("todo", "doing", "review")
        for item in offered
    )
    dropped = send(
        bob,
        "post",
        SHEET,
        {"record": offered[0]["id"], "start": "2026-10-07T10:00"},
    )
    assert dropped.status_code == 201, dropped.content
    assert dropped.json()["startsAt"] == "2026-10-07T10:00"
    assert dropped.json()["endsAt"] == "2026-10-07T11:00"
    logged = TimeEntry.objects.get(pk=dropped.json()["id"])
    assert logged.user == user_bob and str(logged.hours) == "1.00"

    # Moved, then made longer, on the hours.
    MOVE = f"/api/planner/timeentry/{logged.pk}/calendars/timesheet/"
    assert send(
        bob, "patch", MOVE, {"start": "2026-10-08T14:00"}
    ).status_code == 200
    assert send(
        bob, "patch", MOVE, {"end": "2026-10-08T16:30"}
    ).status_code == 200
    logged.refresh_from_db()
    assert str(logged.hours) == "2.50", logged.hours
    assert send(
        bob, "patch", MOVE, {"end": "2026-10-08T13:00"}
    ).status_code == 400

    # A reader neither drops nor moves.
    assert viewer.get(SHEET + "sources/").status_code == 200
    assert send(
        viewer, "post", SHEET,
        {"record": offered[0]["id"], "start": "2026-10-07T10:00"},
    ).status_code == 403

    # Verifying is a lead's.
    def run(client, name, ids):
        return send(
            client, "post", "/api/planner/timeentry/actions/",
            {"action": name, "ids": ids},
        )

    assert run(bob, "verify", [logged.pk]).status_code in (400, 403)
    assert run(signed_in("alice"), "verify", [logged.pk]).status_code == 200
    logged.refresh_from_db()
    assert logged.verified and logged.verified_by.username == "alice"

    print("ok")
    """)


def manage(*args: str) -> subprocess.CompletedProcess:
    env = {
        name: value
        for name, value in os.environ.items()
        if name != "DJANGO_SETTINGS_MODULE"
    }

    return subprocess.run(
        [sys.executable, *args],
        cwd=LOGBOOK,
        env=env,
        capture_output=True,
        text=True,
        timeout=180,
    )


def test_notes_tasks_and_time():
    result = manage("-c", SCENARIO)

    assert result.returncode == 0, result.stdout + result.stderr
    assert result.stdout.strip().endswith("ok")


def test_the_checks_find_nothing():
    result = manage("manage.py", "check", "--fail-level", "INFO")

    assert result.returncode == 0, result.stderr
    assert "no issues" in result.stdout


def test_the_migrations_are_up_to_date():
    result = manage("manage.py", "makemigrations", "--check", "--dry-run")

    assert result.returncode == 0, result.stdout + result.stderr


#: The old planner's export (a sample written by export_planner.py from
#: an old planner, with its API's answers), imported twice, then the
#: key figures, the imputation export and the workdays page.
LEGACY = textwrap.dedent("""
    import io
    import json
    import os

    os.environ["DJANGO_SETTINGS_MODULE"] = "logsite.settings"

    import django
    from django.conf import settings

    django.setup()
    settings.ALLOWED_HOSTS = ["testserver"]

    from django.contrib.auth import get_user_model
    from django.core.management import call_command
    from django.db import connection
    from django.test import Client
    from django.test.utils import setup_test_environment

    setup_test_environment()
    connection.creation.create_test_db(verbosity=0)
    owner = get_user_model().objects.create_superuser("bastien", "", "x")

    from planner.models import (
        ImputationShare, KpiRecord, Project, Task, TimeEntry,
    )

    EXPORT = "migration/sample_export.json"
    data = json.load(open(EXPORT, encoding="utf-8"))

    # A dry run keeps nothing.
    out = io.StringIO()
    call_command(
        "import_planner", EXPORT, "--user", "bastien", "--dry-run",
        stdout=out,
    )
    assert "Dry run" in out.getvalue()
    assert not Task.objects.exists()

    # Twice: nothing doubled, and the key figures as the old API said.
    for _ in range(2):
        out = io.StringIO()
        call_command("import_planner", EXPORT, "--user", "bastien", stdout=out)
        report = out.getvalue()
        assert "differ" not in report, report
        assert "periods identical to the old API" in report, report

    assert Task.objects.count() == len(data["taches"])
    assert TimeEntry.objects.count() == len(data["emargements"])
    assert Project.objects.count() == len(data["activites"])
    assert ImputationShare.objects.count() == len(data["matrice"])
    assert KpiRecord.objects.count() == (
        len(data["kpi_mois"]) + len(data["kpi_annee"])
    )
    seconds = sum(
        (e.ended_at - e.started_at).total_seconds()
        for e in TimeEntry.objects.all()
    )
    assert seconds == data["totaux"]["secondes"], seconds
    assert all(e.user == owner and e.verified for e in TimeEntry.objects.all())
    tickets = Task.objects.filter(kind="FNC")
    assert tickets.count() == sum(t["non_conformite"] for t in data["taches"])
    assert not tickets.filter(criticality="").exists()

    # The API, at its old address, without signing in.
    client = Client()
    for label, old in data["kpi_reponses"]["mois"].items():
        new = client.get("/planner/api/kpi/", {"date": label}).json()
        assert new == old, (label, new, old)
    for label, old in data["kpi_reponses"]["annee"].items():
        new = client.get("/planner/api/kpi/year/", {"date": label}).json()
        assert new == old, (label, new, old)
    assert client.get("/planner/api/kpi/", {"date": "x"}).json() == {
        "erreur": "Date passée en argument incohérente"
    }

    # With a token set, only who sends it.
    settings.LOGBOOK_KPI_TOKEN = "s3cret"
    assert client.get("/planner/api/kpi/").status_code == 401
    assert client.get(
        "/planner/api/kpi/", HTTP_AUTHORIZATION="Token s3cret"
    ).status_code == 200
    settings.LOGBOOK_KPI_TOKEN = ""

    # The imputation export: the time on 7.4-hour days, in Excel.
    client.force_login(owner)
    ids = list(TimeEntry.objects.values_list("pk", flat=True)[:10])
    answer = client.post(
        "/api/planner/timeentry/actions/",
        json.dumps({"action": "export_imputations", "ids": ids}),
        content_type="application/json",
    ).json()
    response = client.get(answer["redirect"])
    assert response.status_code == 200
    assert "repartition_imputations_" in response["Content-Disposition"]

    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(response.content))
    rows = list(book["Repartition"].iter_rows(values_only=True))[1:]
    summary = dict(book["Resume"].iter_rows(values_only=True))
    hours = sum(
        (e.ended_at - e.started_at).total_seconds()
        for e in TimeEntry.objects.filter(pk__in=ids)
    ) / 3600
    assert summary["Temps selectionne sur base 8 h"] == hours
    assert abs(sum(r[1] for r in rows) - hours * 7.4 / 8) < 1e-6
    assert all(r[1] <= 7.4 + 1e-9 for r in rows)

    # No matrix, no file: said, and back to the list.
    ImputationShare.objects.all().delete()
    response = client.get(answer["redirect"])
    assert response.status_code == 302

    # The workdays page.
    assert client.get(
        "/planner/timeentry/workdays/", {"month": "2025-10"}
    ).status_code == 200

    print("ok")
    """)


def test_the_old_planner_imported():
    result = manage("-c", LEGACY)

    assert result.returncode == 0, result.stdout + result.stderr
    assert result.stdout.strip().endswith("ok")
