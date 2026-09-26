"""Exports: every row the table shows, in a file."""

from __future__ import annotations

import csv
import io

import pytest
from openpyxl import load_workbook

pytestmark = pytest.mark.django_db

TICKETS = "/api/example/ticket/"


def content(response) -> bytes:
    if response.streaming:
        return b"".join(response.streaming_content)

    return response.content


def sheet_rows(response) -> list[tuple]:
    workbook = load_workbook(io.BytesIO(content(response)))

    return list(workbook.active.iter_rows(values_only=True))


def by_reference(rows: list) -> tuple[list, dict]:
    header = list(rows[0])
    reference = header.index("Reference")

    return header, {row[reference]: row for row in rows[1:]}


class TestSpreadsheet:
    def test_relations_are_written_as_their_label(
        self,
        worker_client,
        support_desk,
    ):
        response = worker_client.get(f"{TICKETS}export/")
        header, rows = by_reference(sheet_rows(response))

        assert response.status_code == 200
        assert rows["SD-1"][header.index("Team")] == "Front office"
        assert rows["SD-1"][header.index("Assignee")] == "Camille Rousseau"
        assert rows["SD-1"][header.index("Tags")] == "regression, release"
        # A choice drawn as a tag is written as its label.
        assert rows["SD-1"][header.index("Status")] == "Open"
        assert rows["SD-3"][header.index("Assignee")] is None

    def test_every_row_matching_the_search_is_written(
        self,
        worker_client,
        support_desk,
    ):
        response = worker_client.get(
            f"{TICKETS}export/",
            {"search[value]": "invoice"},
        )
        _, rows = by_reference(sheet_rows(response))

        assert sorted(rows) == ["SD-2", "SD-3"]

    def test_the_file_is_named_after_the_model(
        self,
        worker_client,
        support_desk,
    ):
        response = worker_client.get(f"{TICKETS}export/")

        assert "ticket-" in response["Content-Disposition"]
        assert response["Content-Disposition"].endswith('.xlsx"')


class TestCsv:
    def test_relations_are_written_as_their_label(
        self,
        worker_client,
        support_desk,
    ):
        response = worker_client.get(f"{TICKETS}export-csv/")
        text = content(response).decode("utf-8-sig")
        # Semicolons: what Excel expects outside the English locales.
        reader = csv.reader(io.StringIO(text), delimiter=";")
        header, rows = by_reference(list(reader))

        assert response.status_code == 200
        assert rows["SD-2"][header.index("Team")] == "Infrastructure"
        assert rows["SD-2"][header.index("Tags")] == "billing"
