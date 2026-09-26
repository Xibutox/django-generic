"""Exports: everything the filters return, and nothing else."""

from __future__ import annotations

import datetime
import io
import json

import pytest
from django.test import override_settings
from openpyxl import load_workbook

from tests.factories import BookFactory

pytestmark = pytest.mark.django_db

CSV_URL = "/api/books/export-csv/"
XLSX_URL = "/api/books/export/"


def read_csv(response) -> list[list[str]]:
    content = b"".join(response.streaming_content).decode("utf-8")
    # Strip the BOM Excel needs.
    content = content.lstrip("\ufeff")

    return [line.split(";") for line in content.splitlines() if line.strip()]


class TestCsvExport:
    def test_header_uses_the_column_titles(self, api_client, library):
        rows = read_csv(api_client.get(CSV_URL))

        assert rows[0][:3] == ["Title", "Author", "Genre"]

    def test_every_row_is_exported_ignoring_pagination(
        self,
        api_client,
        db,
    ):
        for _ in range(25):
            BookFactory()

        rows = read_csv(api_client.get(CSV_URL, {"length": "5"}))

        # 25 rows plus the header, despite a page size of 5.
        assert len(rows) == 26

    def test_the_export_honours_the_active_filters(
        self,
        api_client,
        library,
    ):
        response = api_client.get(
            CSV_URL,
            {
                "advanced_filters": json.dumps(
                    {
                        "genre": {
                            "operator": "include",
                            "value": ["essay"],
                        }
                    }
                )
            },
        )
        rows = read_csv(response)

        assert len(rows) == 2
        assert rows[1][0] == "Essays"

    def test_column_selection_narrows_the_export(
        self,
        api_client,
        library,
    ):
        rows = read_csv(api_client.get(CSV_URL, {"columns": "title,pages"}))

        assert rows[0] == ["Title", "Pages"]

    def test_an_unknown_column_selection_falls_back_to_everything(
        self,
        api_client,
        library,
    ):
        rows = read_csv(api_client.get(CSV_URL, {"columns": "nonexistent"}))

        assert len(rows[0]) > 2

    def test_a_computed_column_is_exported(self, api_client, library):
        rows = read_csv(api_client.get(CSV_URL, {"columns": "title,slug"}))
        values = {row[0]: row[1] for row in rows[1:]}

        assert values["Emma"] == "emma"

    def test_null_values_become_empty_cells(self, api_client, library):
        rows = read_csv(api_client.get(CSV_URL, {"columns": "title,rating"}))
        values = {row[0]: row[1] for row in rows[1:]}

        assert values["Essays"] == ""

    def test_the_response_is_an_attachment(self, api_client, library):
        response = api_client.get(CSV_URL)

        assert "attachment" in response["Content-Disposition"]
        assert ".csv" in response["Content-Disposition"]

    def test_an_oversized_export_is_refused(self, api_client, library):
        with override_settings(GENERIC={"EXPORT_MAX_ROWS": 2}):
            response = api_client.get(CSV_URL)

        assert response.status_code == 400
        assert "limit" in str(response.data).lower()


class TestExcelExport:
    def test_the_workbook_carries_the_rows(self, api_client, library):
        response = api_client.get(XLSX_URL)

        assert response.status_code == 200

        content = b"".join(response.streaming_content)
        workbook = load_workbook(io.BytesIO(content))
        sheet = workbook.active
        rows = list(sheet.values)

        assert rows[0][0] == "Title"
        assert {row[0] for row in rows[1:]} == {
            "Emma",
            "Persuasion",
            "Essays",
        }

    def test_dates_stay_dates_rather_than_strings(
        self,
        api_client,
        library,
    ):
        response = api_client.get(
            XLSX_URL,
            {"columns": "title,published_on"},
        )
        content = b"".join(response.streaming_content)
        sheet = load_workbook(io.BytesIO(content)).active
        values = {row[0]: row[1] for row in list(sheet.values)[1:]}

        assert isinstance(values["Emma"], datetime.datetime)
        assert values["Emma"].date() == datetime.date(1815, 12, 23)

    def test_datetimes_are_made_naive_in_local_time(
        self,
        api_client,
        library,
    ):
        response = api_client.get(
            XLSX_URL,
            {"columns": "title,released_at"},
        )
        content = b"".join(response.streaming_content)
        sheet = load_workbook(io.BytesIO(content)).active
        values = {row[0]: row[1] for row in list(sheet.values)[1:]}
        released = values["Emma"]

        assert released.tzinfo is None
        # 08:30 UTC is 09:30 in Europe/Paris.
        assert released.hour == 9

    def test_the_export_honours_the_filters(self, api_client, library):
        response = api_client.get(
            XLSX_URL,
            {
                "advanced_filters": json.dumps(
                    {
                        "title": {
                            "operator": "contains",
                            "value": "Emma",
                        }
                    }
                )
            },
        )
        content = b"".join(response.streaming_content)
        sheet = load_workbook(io.BytesIO(content)).active

        assert len(list(sheet.values)) == 2
