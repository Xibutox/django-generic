"""Scheduled mailings: a list, as each reader may see it, by e-mail.

Mail goes to Django's locmem outbox; time is frozen where the schedule
is the subject.
"""

from __future__ import annotations

import datetime
import json
import warnings
import zoneinfo

import pytest
from django.contrib.auth.models import Group, Permission
from django.core import mail
from django.test import override_settings
from freezegun import freeze_time
from openpyxl import load_workbook

from example.models import Ticket
from generic.mailings.dispatch import claim, dispatch
from generic.mailings.models import ScheduledMailing
from generic.mailings.schedule import next_run
from generic.mailings.sending import parameters, send
from tests.factories import UserFactory

pytestmark = pytest.mark.django_db

API = "/api/generic/scheduledmailing/"
OPEN_TICKETS = {
    "columns": ["reference", "title", "status"],
    "order": [["reference", "asc"]],
    "filters": {
        "match": "all",
        "conditions": [
            {"column": "status", "operator": "any_of", "value": ["open"]}
        ],
    },
    "search": "",
}
PARIS = zoneinfo.ZoneInfo("Europe/Paris")


def person(username, *codenames, email=True, **extra):
    user = UserFactory(
        username=username,
        email=f"{username}@example.test" if email else "",
        **extra,
    )
    user.user_permissions.set(
        Permission.objects.filter(codename__in=codenames)
    )

    return user


def mailing(owner, **fields) -> ScheduledMailing:
    values = {
        "name": "Open tickets",
        "table": "site.example.ticket",
        "state": OPEN_TICKETS,
        "owner": owner,
        **fields,
    }
    return ScheduledMailing.objects.create(**values)


def rows_of(message) -> list[list]:
    filename, content, _mimetype = message.attachments[0]
    workbook = load_workbook(__import__("io").BytesIO(content))

    return [list(row) for row in workbook.active.iter_rows(values_only=True)]


class TestTheSchedule:
    def at(self, *args):
        return datetime.datetime(*args, tzinfo=PARIS)

    @override_settings(TIME_ZONE="Europe/Paris")
    def test_every_day_is_tomorrow_once_today_has_passed(self):
        entry = ScheduledMailing(frequency="daily", time=datetime.time(8))

        assert next_run(entry, self.at(2026, 3, 10, 9)) == self.at(
            2026, 3, 11, 8
        )
        assert next_run(entry, self.at(2026, 3, 10, 7)) == self.at(
            2026, 3, 10, 8
        )

    @override_settings(TIME_ZONE="Europe/Paris")
    def test_weekdays_skip_the_weekend(self):
        entry = ScheduledMailing(frequency="weekdays", time=datetime.time(8))

        # A Friday afternoon: next is Monday morning.
        assert next_run(entry, self.at(2026, 9, 25, 15)) == self.at(
            2026, 9, 28, 8
        )

    @override_settings(TIME_ZONE="Europe/Paris")
    def test_weekly_on_its_day(self):
        entry = ScheduledMailing(
            frequency="weekly", time=datetime.time(8), weekday=2
        )

        assert next_run(entry, self.at(2026, 9, 25, 15)).date() == (
            datetime.date(2026, 9, 30)
        )

    @override_settings(TIME_ZONE="Europe/Paris")
    def test_monthly_across_the_end_of_the_year(self):
        entry = ScheduledMailing(
            frequency="monthly", time=datetime.time(8), day_of_month=15
        )

        assert next_run(entry, self.at(2026, 12, 20)) == self.at(
            2027, 1, 15, 8
        )

    @override_settings(TIME_ZONE="Europe/Paris")
    def test_the_hour_stays_local_across_summer_time(self):
        entry = ScheduledMailing(frequency="daily", time=datetime.time(8))

        # Summer time starts on 29 March 2026: 08:00 is 07:00 UTC the
        # day before, 06:00 UTC the day after.
        before = next_run(entry, self.at(2026, 3, 27, 9))
        after = next_run(entry, self.at(2026, 3, 29, 9))

        assert before.astimezone(datetime.timezone.utc).hour == 7
        assert after.astimezone(datetime.timezone.utc).hour == 6


class TestSending:
    def test_each_reader_gets_the_rows_they_may_see(self, support_desk):
        """The owner sees every ticket; a reader whose queryset is
        narrowed by a permission they lack gets nothing at all."""
        owner = person("owner", "view_ticket", "add_scheduledmailing")
        stranger = person("stranger")
        entry = mailing(owner)
        entry.users.add(stranger)

        outcome = send(entry)

        assert outcome.sent == 1
        assert outcome.refused == 1
        assert [message.to for message in mail.outbox] == [
            ["owner@example.test"]
        ]
        assert rows_of(mail.outbox[0]) == [
            ["Reference", "Title", "Status"],
            ["SD-1", "Login fails after password reset", "Open"],
        ]

    def test_the_rows_follow_the_readers_own_scope(
        self, support_desk, monkeypatch
    ):
        from example.resources import TicketResource

        owner = person("owner", "view_ticket", "add_scheduledmailing")
        narrow = person("narrow", "view_ticket")
        base = TicketResource.get_queryset

        def only_their_team(self, request):
            queryset = base(self, request)

            if request.user.username == "narrow":
                return queryset.filter(team__code="INFRA")

            return queryset

        monkeypatch.setattr(TicketResource, "get_queryset", only_their_team)
        entry = mailing(owner, state={}, send_when_empty=True)
        entry.users.add(narrow)

        send(entry)

        by_reader = {
            message.to[0]: [row[0] for row in rows_of(message)[1:]]
            for message in mail.outbox
        }
        assert sorted(by_reader["owner@example.test"]) == [
            "SD-1",
            "SD-2",
            "SD-3",
        ]
        assert by_reader["narrow@example.test"] == ["SD-2"]

    def test_an_empty_list_is_not_sent(self, support_desk):
        owner = person("owner", "view_ticket")
        Ticket.objects.update(status="closed")

        outcome = send(mailing(owner))

        assert outcome.empty == 1
        assert mail.outbox == []

    def test_unless_asked(self, support_desk):
        owner = person("owner", "view_ticket")
        Ticket.objects.update(status="closed")

        send(mailing(owner, send_when_empty=True))

        assert len(mail.outbox) == 1

    def test_csv_too(self, support_desk):
        owner = person("owner", "view_ticket")

        send(mailing(owner, format="csv"))

        filename, content, mimetype = mail.outbox[0].attachments[0]
        assert filename.endswith(".csv")
        assert mimetype == "text/csv"
        # A text attachment is kept as text by the mail machinery.
        assert "SD-1" in (
            content.decode() if isinstance(content, bytes) else content
        )

    def test_the_mail_links_the_filtered_list(self, support_desk):
        owner = person("owner", "view_ticket")

        send(mailing(owner))

        assert "/example/ticket/?filters=" in mail.outbox[0].body

    @override_settings(GENERIC={"MAILING_MAX_ATTACHMENT_SIZE": 10})
    def test_a_file_too_large_is_linked_instead(self, support_desk):
        owner = person("owner", "view_ticket")

        send(mailing(owner))

        assert mail.outbox[0].attachments == []
        assert "too large" in mail.outbox[0].body

    def test_groups_inactive_accounts_and_missing_addresses(
        self, support_desk
    ):
        owner = person("owner", "view_ticket")
        team = Group.objects.create(name="Desk")
        member = person("member", "view_ticket")
        gone = person("gone", "view_ticket", is_active=False)
        silent = person("silent", "view_ticket", email=False)
        team.user_set.add(member, gone, silent)
        entry = mailing(owner)
        entry.groups.add(team)
        entry.users.add(member)

        assert [user.username for user in entry.recipients()] == [
            "owner",
            "member",
        ]

    def test_the_mail_speaks_the_readers_language(self, support_desk):
        from generic.accounts.models import UserPreferences

        owner = person("owner", "view_ticket")
        preferences = UserPreferences.for_user(owner)
        preferences.language = "fr"
        preferences.save()

        send(mailing(owner))

        assert mail.outbox[0].body.startswith("1 ligne de Tickets")
        assert rows_of(mail.outbox[0])[0][0] == "Référence"

    def test_a_column_that_is_gone_pauses_and_tells_the_owner(
        self, support_desk
    ):
        from generic.events.models import Notification

        owner = person("owner", "view_ticket")
        state = {
            "filters": {
                "match": "all",
                "conditions": [
                    {"column": "vanished", "operator": "equals", "value": "x"}
                ],
            }
        }
        entry = mailing(owner, state=state)

        send(entry)

        entry.refresh_from_db()
        assert entry.is_active is False
        assert entry.last_error
        assert Notification.objects.filter(user=owner).exists()

    def test_sending_warns_of_nothing(self, support_desk):
        """Django 6.1 deprecates the old mail settings: none are used."""
        owner = person("owner", "view_ticket")

        with warnings.catch_warnings():
            warnings.simplefilter("error")
            send(mailing(owner))

        assert len(mail.outbox) == 1


class TestTheDispatcher:
    def test_it_sends_what_is_due_and_moves_it_on(self, support_desk):
        owner = person("owner", "view_ticket")
        due = mailing(owner, next_run_at="2020-01-01T00:00:00Z")
        later = mailing(
            owner, name="Later", next_run_at="2999-01-01T00:00:00Z"
        )

        assert dispatch() == 1

        due.refresh_from_db()
        later.refresh_from_db()
        assert due.next_run_at.year > 2020
        assert due.last_sent_at is not None
        assert later.last_sent_at is None

    def test_a_mailing_already_claimed_is_not_sent_twice(self, support_desk):
        """The second dispatcher finds it moved on: nothing to send."""
        owner = person("owner", "view_ticket")
        entry = mailing(owner, next_run_at="2020-01-01T00:00:00Z")

        with freeze_time("2026-09-26 10:00:00"):
            from django.utils import timezone

            now = timezone.now()

            assert claim(entry.pk, now) is not None
            assert claim(entry.pk, now) is None

    def test_a_paused_mailing_waits(self, support_desk):
        owner = person("owner", "view_ticket")
        mailing(owner, next_run_at="2020-01-01T00:00:00Z", is_active=False)

        assert dispatch() == 0

    def test_the_task_is_declared(self):
        from generic.tasks import get_task

        assert get_task("generic.send_scheduled_mailings").label

    def test_the_command(self, support_desk):
        from io import StringIO

        from django.core.management import call_command

        owner = person("owner", "view_ticket")
        mailing(owner, next_run_at="2020-01-01T00:00:00Z")
        out = StringIO()

        call_command("send_scheduled_mailings", stdout=out)

        assert "1 mailing" in out.getvalue()


class TestTheScreens:
    def test_a_list_offers_it_to_whoever_may_add_one(
        self, client, support_desk
    ):
        user = person("maker", "view_ticket", "add_scheduledmailing")
        client.force_login(user)

        response = client.get("/example/ticket/")
        options = response.context["table_config"]["options"]

        assert options["mailingUrl"] == "/generic/scheduledmailing/add/"

    def test_and_to_nobody_else(self, client, support_desk):
        client.force_login(person("reader", "view_ticket"))

        response = client.get("/example/ticket/")

        assert response.context["table_config"]["options"]["mailingUrl"] == ""

    def test_a_resource_may_refuse_to_be_mailed(
        self, client, support_desk, monkeypatch
    ):
        from example.resources import TicketResource

        monkeypatch.setattr(TicketResource, "mailing", False)
        user = person("maker", "view_ticket", "add_scheduledmailing")
        client.force_login(user)

        response = client.post(
            API,
            {"name": "x", "table": "site.example.ticket", "state": {}},
            content_type="application/json",
        )

        assert response.status_code == 400
        assert "table" in response.json()

    def test_the_add_form_takes_the_table_from_the_address(
        self, client, support_desk
    ):
        client.force_login(
            person("maker", "view_ticket", "add_scheduledmailing")
        )

        response = client.get(
            "/generic/scheduledmailing/add/",
            {
                "table": "site.example.ticket",
                "state": json.dumps(OPEN_TICKETS),
            },
        )

        initial = response.context["form_config"]["initial"]
        assert initial["state"] == OPEN_TICKETS

    def test_creating_one_schedules_it_for_its_owner(
        self, client, support_desk
    ):
        maker = person("maker", "view_ticket", "add_scheduledmailing")
        client.force_login(maker)

        response = client.post(
            API,
            {
                "name": "Open tickets",
                "table": "site.example.ticket",
                "state": OPEN_TICKETS,
                "frequency": "weekdays",
                "time": "08:00",
            },
            content_type="application/json",
        )

        assert response.status_code == 201, response.content
        entry = ScheduledMailing.objects.get()
        assert entry.owner == maker
        assert entry.next_run_at is not None

    def test_a_list_the_owner_may_not_see_is_refused(
        self, client, support_desk
    ):
        client.force_login(person("maker", "add_scheduledmailing"))

        response = client.post(
            API,
            {"name": "x", "table": "site.example.ticket", "state": {}},
            content_type="application/json",
        )

        assert response.status_code == 400

    def test_a_day_of_the_month_past_28_is_refused(self, client, support_desk):
        client.force_login(
            person("maker", "view_ticket", "add_scheduledmailing")
        )

        response = client.post(
            API,
            {
                "name": "x",
                "table": "site.example.ticket",
                "frequency": "monthly",
                "day_of_month": 31,
            },
            content_type="application/json",
        )

        assert response.status_code == 400
        assert "day_of_month" in response.json()

    def test_people_see_their_own_mailings_only(self, client, support_desk):
        mine = person("mine", "view_ticket", "add_scheduledmailing")
        theirs = person("theirs", "view_ticket", "add_scheduledmailing")
        mailing(mine, name="Mine")
        other = mailing(theirs, name="Theirs")
        client.force_login(mine)

        rows = client.get(API, {"draw": 1}).json()["data"]

        assert [row["name"] for row in rows] == ["Mine"]
        assert client.get(f"{API}{other.pk}/").status_code == 404

    def test_the_view_permission_sees_everyones(self, client, support_desk):
        mailing(person("a", "view_ticket"), name="A")
        mailing(person("b", "view_ticket"), name="B")
        client.force_login(person("boss", "view_scheduledmailing"))

        rows = client.get(API, {"draw": 1}).json()["data"]

        assert sorted(row["name"] for row in rows) == ["A", "B"]

    def test_send_now(self, client, support_desk):
        owner = person("owner", "view_ticket", "add_scheduledmailing")
        entry = mailing(owner)
        client.force_login(owner)

        response = client.post(
            f"{API}actions/",
            {"action": "send_now", "ids": [entry.pk]},
            content_type="application/json",
        )

        assert response.status_code == 200, response.content
        assert len(mail.outbox) == 1

    def test_pause_and_resume(self, client, support_desk):
        owner = person("owner", "view_ticket", "add_scheduledmailing")
        entry = mailing(owner)
        client.force_login(owner)

        for name, active in (("pause", False), ("resume", True)):
            client.post(
                f"{API}actions/",
                {"action": name, "ids": [entry.pk]},
                content_type="application/json",
            )
            entry.refresh_from_db()

            assert entry.is_active is active

    @override_settings(GENERIC={"SHOW_MAILINGS": False})
    def test_switched_off_nothing_is_offered(self, client, support_desk):
        client.force_login(
            person("maker", "view_ticket", "add_scheduledmailing")
        )

        response = client.get("/example/ticket/")

        assert response.context["table_config"]["options"]["mailingUrl"] == ""


def test_a_state_becomes_the_tables_own_parameters():
    params = parameters(
        {
            "columns": ["reference", "title"],
            "order": [["due_on", "desc"], ["reference", "asc"]],
            "filters": OPEN_TICKETS["filters"],
            "search": " invoice ",
            "pageLength": 25,
        }
    )

    assert params == {
        "columns": "reference,title",
        "ordering": "-due_on,reference",
        "filters": json.dumps(OPEN_TICKETS["filters"]),
        "search": "invoice",
    }


def test_the_framework_task_alone_does_not_bring_the_task_pages(monkeypatch):
    """A project that declared no task keeps its navigation as it was."""
    from generic.tasks import resources
    from generic.tasks.registry import registry

    monkeypatch.setattr(
        registry, "names", lambda: ["generic.send_scheduled_mailings"]
    )
    monkeypatch.setattr(resources, "beat_installed", lambda: False)

    assert resources.tasks_are_offered() is False
