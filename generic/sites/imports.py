"""Imports: a spreadsheet read back into records.

The reverse of the exports, declared on the resource and off unless it
is::

    @register(Ticket)
    class TicketResource(ModelResource):
        imports = Import(
            fields=("reference", "title", "customer", "status", "tags"),
            key="reference",                  # update rows it finds
            lookups={"customer": "code"},     # how a cell finds a record
        )

A file goes through five steps, every one of them the same for the
preview and for the real thing:

1. **read** - ``.xlsx`` through openpyxl (values, never formulas run),
   ``.csv`` through the standard library, limits checked as it reads;
2. **map** - each header matched to a column, ignoring case and accents,
   on its name, the title the export writes, or the field's name - so an
   export imports back as it is;
3. **convert** - each cell into what the form would send: choices by
   value or label, yes/no, dates and decimals in the reader's language,
   relations by label among the records the reader may see;
4. **validate** - each row through the resource's own form serializer,
   with its unique checks; updates are partial, an empty cell keeps the
   stored value;
5. **write** - all or nothing, in one transaction, as the importer
   (``acting_as``), with one live ``bulk`` event.

The preview is the same run, rolled back: unique and database
constraints are checked for real, and nothing is kept between the
preview and the confirmation - the browser sends the file again.
"""

from __future__ import annotations

import csv
import dataclasses
import datetime
import decimal
import io
import re
from typing import Any, Callable, Iterable, Iterator, Sequence

from django.core.exceptions import (
    FieldDoesNotExist,
    ImproperlyConfigured,
    ValidationError,
)
from django.db import IntegrityError, models, router, transaction
from django.utils import formats, timezone
from django.utils.dateparse import parse_date, parse_datetime, parse_time
from django.utils.encoding import force_str
from django.utils.translation import gettext
from rest_framework.exceptions import ValidationError as ApiValidationError

from generic.conf import generic_settings
from generic.search import fold, text_lookup

#: What a declaration may ask for.
MODES = ("create", "update", "create_update")

#: The formats read, by extension.
FORMATS = (".xlsx", ".csv")

#: Beyond this many columns a file is not a table of records.
MAX_COLUMNS = 200

#: Errors listed, at most: past that, the file needs another look.
MAX_ERRORS = 200

#: What a cell says to mean yes, and no - in the two languages shipped,
#: and the spellings a spreadsheet writes.
TRUE_WORDS = frozenset(
    {"1", "x", "yes", "y", "true", "vrai", "oui", "o", "on", "✓"}
)
FALSE_WORDS = frozenset({"0", "no", "n", "false", "faux", "non", "off", ""})

#: Several labels in one cell: what the export writes between them.
SEPARATOR = re.compile(r"\s*[,;\n]\s*")


class ImportRefused(Exception):
    """A file that cannot be imported at all: the message says why."""

    def __init__(self, message: Any) -> None:
        super().__init__(force_str(message))
        self.message = force_str(message)


@dataclasses.dataclass(frozen=True)
class Import:
    """How a resource's records may be imported.

    ``fields``
        The fields a file may fill, by name. Default: the form's
        editable fields.
    ``key``
        A unique field matching a row to an existing record, which the
        row then updates. None: every row is a new record.
    ``mode``
        ``create``, ``update`` or ``create_update``. Default: the last
        with a ``key``, the first without.
    ``lookups``
        ``{relation: field}``: the field of the related model a cell is
        matched against. Default: the related model's naming field.
    ``defaults``
        ``{field: value}``, or ``callable(request)`` returning one, for
        new records: values the file does not carry.
    ``permission``
        Needed on top of the resource's add or change permission: a
        permission string, or ``callable(user)``.
    ``max_rows``
        Default ``GENERIC["IMPORT_MAX_ROWS"]``.
    ``description``
        Shown on the import page: what a row is, what to leave empty.
    """

    fields: Sequence[str] | None = None
    key: str | None = None
    mode: str | None = None
    lookups: dict[str, str] = dataclasses.field(default_factory=dict)
    defaults: dict[str, Any] | Callable[[Any], dict[str, Any]] | None = None
    permission: str | Callable[[Any], bool] | None = None
    max_rows: int | None = None
    description: Any = ""

    @property
    def resolved_mode(self) -> str:
        if self.mode:
            return self.mode

        return "create_update" if self.key else "create"

    @property
    def creates(self) -> bool:
        return self.resolved_mode in ("create", "create_update")

    @property
    def updates(self) -> bool:
        return self.resolved_mode in ("update", "create_update")


def declaration_of(resource: Any) -> Import | None:
    """The resource's ``imports``, read: ``True`` is ``Import()``."""
    value = getattr(resource, "imports", None)

    if value is True:
        return Import()

    if value in (None, False):
        return None

    if not isinstance(value, Import):
        raise ImproperlyConfigured(
            f"{type(resource).__name__}.imports must be an Import(...), "
            f"True or None, not {value!r}."
        )

    return value


# ---------------------------------------------------------------------
# The declaration, checked
# ---------------------------------------------------------------------


def model_field(model: Any, name: str) -> Any:
    try:
        return model._meta.get_field(name)
    except FieldDoesNotExist:
        return None


def is_importable_field(field: Any) -> bool:
    """A field a file can write: the model's own, editable, not a key."""
    if field is None or field.auto_created and not field.concrete:
        return False

    if getattr(field, "primary_key", False) and isinstance(
        field, models.AutoField
    ):
        return False

    if field.many_to_many:
        return field.concrete and field.editable

    return bool(field.concrete and field.editable)


def check_import(resource: Any) -> None:
    """What the model alone can tell about a declaration, at registration.

    What needs the form serializer - is the field in the form - waits
    for the first use, like a grid's: other resources may not be
    registered yet.
    """
    declaration = declaration_of(resource)

    if declaration is None:
        return

    name = type(resource).__name__
    model = resource.model

    if declaration.resolved_mode not in MODES:
        raise ImproperlyConfigured(
            f"{name}.imports: mode must be one of {', '.join(MODES)}, not "
            f"{declaration.mode!r}."
        )

    if declaration.updates and not declaration.key:
        raise ImproperlyConfigured(
            f"{name}.imports: mode {declaration.resolved_mode!r} updates "
            f"records and needs a key to find them."
        )

    for field_name in declaration.fields or ():
        if not is_importable_field(model_field(model, field_name)):
            raise ImproperlyConfigured(
                f"{name}.imports: {field_name!r} is not an editable field "
                f"of {model.__name__}."
            )

    if declaration.key:
        field = model_field(model, declaration.key)

        if field is None or not (
            getattr(field, "unique", False) or field.primary_key
        ):
            raise ImproperlyConfigured(
                f"{name}.imports: the key {declaration.key!r} must be a "
                f"unique field of {model.__name__}."
            )

    for relation, lookup in declaration.lookups.items():
        field = model_field(model, relation)

        if field is None or not field.is_relation:
            raise ImproperlyConfigured(
                f"{name}.imports: lookups names {relation!r}, which is not "
                f"a relation of {model.__name__}."
            )

        if model_field(field.related_model, lookup.split("__")[0]) is None:
            raise ImproperlyConfigured(
                f"{name}.imports: {field.related_model.__name__} has no "
                f"field {lookup!r} to find a {relation} by."
            )


# ---------------------------------------------------------------------
# Columns
# ---------------------------------------------------------------------


@dataclasses.dataclass
class ImportColumn:
    """One field a file may fill, and how its cells are read."""

    name: str
    title: str
    kind: str
    required: bool
    field: Any
    lookup: str | None = None
    choices: dict[Any, str] = dataclasses.field(default_factory=dict)
    aliases: tuple[str, ...] = ()
    #: Every way a choice may be written - its value, its label in each
    #: language offered - simplified, to its value.
    spellings: dict[str, Any] = dataclasses.field(default_factory=dict)

    def describe(self) -> dict[str, Any]:
        entry: dict[str, Any] = {
            "name": self.name,
            "title": self.title,
            "type": self.kind,
            "required": self.required,
        }

        if self.choices:
            entry["choices"] = [
                {"value": value, "label": label}
                for value, label in self.choices.items()
            ]

        if self.lookup:
            entry["lookup"] = self.lookup

        return entry


def kind_of(field: Any) -> str:
    if field.many_to_many:
        return "relations"

    if field.is_relation:
        return "relation"

    if getattr(field, "choices", None):
        return "choice"

    if isinstance(field, models.BooleanField):
        return "boolean"

    if isinstance(field, models.DateTimeField):
        return "datetime"

    if isinstance(field, models.DateField):
        return "date"

    if isinstance(field, models.TimeField):
        return "time"

    if isinstance(field, models.DecimalField):
        return "decimal"

    if isinstance(field, models.FloatField):
        return "float"

    if isinstance(field, (models.IntegerField, models.AutoField)):
        return "integer"

    return "text"


def simplify(text: Any) -> str:
    """A header as it is compared: no accents, case, stars or colons."""
    value = fold(force_str(text or "")).casefold()
    value = re.sub(r"[\s_]+", " ", value).strip()

    return value.rstrip("*: ").strip()


def choice_spellings(field: Any) -> dict[str, Any]:
    """A choice by its value, or its label in any language offered.

    A file written in French names *Ouvert* where the English page says
    *Open*; whoever imports it may read either.
    """
    from django.conf import settings
    from django.utils import translation

    spellings: dict[str, Any] = {}
    languages = [code for code, _name in getattr(settings, "LANGUAGES", ())]

    for value, _label in field.flatchoices:
        if value not in (None, ""):
            spellings.setdefault(simplify(value), value)

    for code in [translation.get_language(), *languages]:
        if not code:
            continue

        with translation.override(code):
            for value, label in field.flatchoices:
                if value not in (None, ""):
                    spellings.setdefault(simplify(force_str(label)), value)

    return spellings


def label_field_of(model: Any) -> str | None:
    from generic.sites.serializers import label_field_for

    return label_field_for(model)


# ---------------------------------------------------------------------
# Reading a file
# ---------------------------------------------------------------------


@dataclasses.dataclass
class Sheet:
    """A file read: its headers, and its rows with their numbers."""

    headers: list[str]
    rows: list[tuple[int, list[Any]]]


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def read_file(upload: Any, max_rows: int) -> Sheet:
    """The first sheet of a workbook, or a CSV, as headers and rows."""
    name = force_str(getattr(upload, "name", "") or "").lower()
    extension = name[name.rfind(".") :] if "." in name else ""

    if extension not in FORMATS:
        raise ImportRefused(
            gettext("Choose an Excel (.xlsx) or CSV (.csv) file.")
        )

    limit = generic_settings.IMPORT_MAX_FILE_SIZE

    if limit and (getattr(upload, "size", 0) or 0) > limit:
        raise ImportRefused(
            gettext("The file is larger than %(size)s MB.")
            % {"size": round(limit / 1024 / 1024, 1)}
        )

    rows = read_xlsx(upload) if extension == ".xlsx" else read_csv(upload)

    headers: list[str] | None = None
    found: list[tuple[int, list[Any]]] = []

    for number, values in rows:
        if len(values) > MAX_COLUMNS:
            values = values[:MAX_COLUMNS]

        if all(is_blank(value) for value in values):
            continue

        if headers is None:
            headers = [force_str(value or "").strip() for value in values]
            continue

        if len(found) >= max_rows:
            raise ImportRefused(
                gettext(
                    "The file has more than %(count)s rows. Split it and "
                    "import each part."
                )
                % {"count": max_rows}
            )

        found.append((number, list(values)))

    if headers is None:
        raise ImportRefused(gettext("The file is empty."))

    return Sheet(headers=headers, rows=found)


def read_xlsx(upload: Any) -> Iterator[tuple[int, list[Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ImportRefused(
            gettext(
                "Excel files cannot be read here. Save the sheet as CSV "
                "and import that."
            )
        )

    upload.seek(0)

    try:
        # Values only: a formula gives what it last computed, and is
        # never run here.
        workbook = load_workbook(upload, read_only=True, data_only=True)
    except Exception:
        raise ImportRefused(gettext("This is not an Excel file it can read."))

    try:
        sheet = workbook.worksheets[0]

        for number, values in enumerate(
            sheet.iter_rows(values_only=True), start=1
        ):
            yield number, list(values)
    finally:
        workbook.close()


def read_csv(upload: Any) -> Iterator[tuple[int, list[Any]]]:
    upload.seek(0)
    raw = upload.read()

    try:
        text = raw.decode("utf-8-sig")
    except UnicodeDecodeError:
        # What Excel writes a CSV in, on a Western Windows.
        text = raw.decode("cp1252", errors="replace")

    sample = text[:8192]

    try:
        dialect: Any = csv.Sniffer().sniff(sample, delimiters=";,\t")
    except csv.Error:
        dialect = csv.excel
        dialect.delimiter = ";" if sample.count(";") else ","

    for number, values in enumerate(
        csv.reader(io.StringIO(text), dialect), start=1
    ):
        yield number, values


# ---------------------------------------------------------------------
# The run
# ---------------------------------------------------------------------


class RowError(Exception):
    """A cell or a row refused: ``column`` None means the whole row."""

    def __init__(self, message: Any, column: str | None = None) -> None:
        super().__init__(force_str(message))
        self.message = force_str(message)
        self.column = column


class Importer:
    """One resource's import, for one request."""

    def __init__(self, resource: Any, request: Any) -> None:
        self.resource = resource
        self.request = request
        self.declaration = declaration_of(resource)

        if self.declaration is None:
            raise ImproperlyConfigured(
                f"{type(resource).__name__} declares no imports."
            )

        self._lookups: dict[tuple[str, str], Any] = {}

    # -- who may ---------------------------------------------------------

    def is_allowed(self) -> bool:
        request = self.request
        resource = self.resource
        declaration = self.declaration

        may_write = (
            declaration.creates and resource.has_add_permission(request)
        ) or (declaration.updates and resource.has_change_permission(request))

        if not may_write:
            return False

        permission = declaration.permission

        if permission is None:
            return True

        if callable(permission):
            return bool(permission(request.user))

        return request.user.has_perm(permission)

    # -- columns ---------------------------------------------------------

    def get_field_names(self) -> list[str]:
        declared = self.declaration.fields

        if declared is not None:
            names = list(declared)
        else:
            readonly = set(self.resource.get_readonly_fields(self.request))
            names = [
                name
                for name in self.resource.get_fields(self.request)
                if name not in readonly
                and is_importable_field(model_field(self.resource.model, name))
            ]

        key = self.declaration.key

        if key and key not in names:
            names.insert(0, key)

        return names

    def get_serializer_fields(self) -> dict[str, Any]:
        serializer = self.resource.get_form_serializer_class()(
            context={"request": self.request}
        )

        return serializer.fields

    def get_columns(self) -> list[ImportColumn]:
        if hasattr(self, "_columns"):
            return self._columns

        model = self.resource.model
        writable = self.get_serializer_fields()
        titles = self.get_export_titles()
        overrides = self.resource.form_overrides or {}
        columns = []

        for name in self.get_field_names():
            field = model_field(model, name)
            serializer_field = writable.get(name)
            is_key = name == self.declaration.key

            if not is_key and (
                serializer_field is None or serializer_field.read_only
            ):
                raise ImproperlyConfigured(
                    f"{type(self.resource).__name__}.imports: {name!r} is "
                    f"not a writable field of its form."
                )

            kind = kind_of(field)
            title = force_str(
                overrides.get(name, {}).get("label")
                or getattr(serializer_field, "label", None)
                or field.verbose_name
            )
            lookup = None

            if field.is_relation:
                lookup = self.declaration.lookups.get(name) or (
                    label_field_of(field.related_model)
                )

                if lookup is None:
                    lookup = "pk"

            choices = (
                {
                    value: force_str(label)
                    for value, label in field.flatchoices
                    if value not in (None, "")
                }
                if kind == "choice"
                else {}
            )
            spellings = choice_spellings(field) if kind == "choice" else {}
            aliases = {
                simplify(name),
                simplify(title),
                simplify(field.verbose_name),
            }

            if name in titles:
                aliases.add(simplify(titles[name]))

            columns.append(
                ImportColumn(
                    name=name,
                    title=title[:1].upper() + title[1:],
                    kind=kind,
                    required=bool(
                        serializer_field is not None
                        and serializer_field.required
                    ),
                    field=field,
                    lookup=lookup,
                    choices=choices,
                    aliases=tuple(sorted(alias for alias in aliases if alias)),
                    spellings=spellings,
                )
            )

        self._columns = columns

        return columns

    def get_export_titles(self) -> dict[str, str]:
        """What the export writes above each column, by column name."""
        try:
            serializer = self.resource.get_table_serializer_class()
            columns = serializer.get_export_columns()
        except Exception:  # pragma: no cover - a table that cannot say
            return {}

        return {
            column["data"]: force_str(column.get("title", ""))
            for column in columns
        }

    def describe(self) -> dict[str, Any]:
        """Everything the import page needs to know, before a file."""
        declaration = self.declaration

        return {
            "columns": [column.describe() for column in self.get_columns()],
            "key": declaration.key,
            "mode": declaration.resolved_mode,
            "creates": declaration.creates
            and self.resource.has_add_permission(self.request),
            "updates": declaration.updates
            and self.resource.has_change_permission(self.request),
            "description": force_str(declaration.description or ""),
            "maxRows": self.max_rows,
            "maxFileSize": generic_settings.IMPORT_MAX_FILE_SIZE,
            "formats": list(FORMATS),
        }

    @property
    def max_rows(self) -> int:
        return int(
            self.declaration.max_rows or generic_settings.IMPORT_MAX_ROWS
        )

    # -- mapping ---------------------------------------------------------

    def guess_mapping(self, headers: Sequence[str]) -> list[str | None]:
        """Each header's column, or None; a column is taken once."""
        by_alias: dict[str, str] = {}

        for column in self.get_columns():
            for alias in column.aliases:
                by_alias.setdefault(alias, column.name)

        taken: set[str] = set()
        mapping: list[str | None] = []

        for header in headers:
            name = by_alias.get(simplify(header))

            if name in taken:
                name = None

            if name:
                taken.add(name)

            mapping.append(name)

        return mapping

    def read_mapping(
        self,
        headers: Sequence[str],
        raw: Any,
    ) -> list[str | None]:
        """The mapping the reader confirmed, checked; else a guess."""
        if raw in (None, ""):
            return self.guess_mapping(headers)

        if not isinstance(raw, list) or len(raw) != len(headers):
            raise ImportRefused(
                gettext("The columns chosen do not match the file.")
            )

        names = {column.name for column in self.get_columns()}
        mapping: list[str | None] = []
        taken: set[str] = set()

        for value in raw:
            name = force_str(value) if value else None

            if name is not None and name not in names:
                raise ImportRefused(
                    gettext("%(column)s cannot be imported.")
                    % {"column": name}
                )

            if name in taken:
                raise ImportRefused(
                    gettext("Two columns of the file go to %(column)s.")
                    % {"column": self.column(name).title}
                )

            if name:
                taken.add(name)

            mapping.append(name)

        return mapping

    def column(self, name: str) -> ImportColumn:
        return next(
            column for column in self.get_columns() if column.name == name
        )

    # -- converting a cell -----------------------------------------------

    def convert(self, column: ImportColumn, value: Any) -> Any:
        """A cell as the form serializer takes it."""
        kind = column.kind

        if isinstance(value, str):
            value = value.strip()

        if kind == "relations":
            return [
                self.find(column, label)
                for label in SEPARATOR.split(force_str(value))
                if label
            ]

        if kind == "relation":
            return self.find(column, value)

        if kind == "choice":
            return self.convert_choice(column, value)

        if kind == "boolean":
            return self.convert_boolean(column, value)

        if kind in ("date", "datetime", "time"):
            return self.convert_moment(column, value)

        if kind in ("integer", "decimal", "float"):
            return self.convert_number(column, value)

        # Text: a whole number typed in a spreadsheet arrives as 12.0.
        if isinstance(value, float) and value.is_integer():
            value = int(value)

        if isinstance(value, (datetime.date, datetime.time)):
            return value.isoformat()

        return force_str(value)

    def convert_choice(self, column: ImportColumn, value: Any) -> Any:
        if isinstance(value, float) and value.is_integer():
            value = int(value)

        wanted = simplify(value)

        if wanted in column.spellings:
            return column.spellings[wanted]

        raise RowError(
            gettext("%(value)s is not one of the choices.")
            % {"value": force_str(value)},
            column.name,
        )

    def convert_boolean(self, column: ImportColumn, value: Any) -> bool:
        if isinstance(value, bool):
            return value

        word = simplify(value)

        if word in TRUE_WORDS:
            return True

        if word in FALSE_WORDS:
            return False

        raise RowError(
            gettext("%(value)s is neither yes nor no.")
            % {"value": force_str(value)},
            column.name,
        )

    def convert_moment(self, column: ImportColumn, value: Any) -> str:
        kind = column.kind
        parsed: Any = None

        if isinstance(value, datetime.datetime):
            parsed = value
        elif isinstance(value, datetime.date):
            parsed = value
        elif isinstance(value, datetime.time):
            parsed = value
        else:
            text = force_str(value)
            parsed = (
                parse_time(text)
                if kind == "time"
                else parse_datetime(text) or parse_date(text)
            )

            if parsed is None:
                parsed = self.parse_localized(kind, text)

        if parsed is None:
            raise RowError(
                gettext("%(value)s is not a date this can read.")
                % {"value": force_str(value)},
                column.name,
            )

        if kind == "date" and isinstance(parsed, datetime.datetime):
            parsed = parsed.date()

        if kind == "datetime":
            if not isinstance(parsed, datetime.datetime):
                parsed = datetime.datetime.combine(parsed, datetime.time())

            if timezone.is_naive(parsed):
                parsed = timezone.make_aware(parsed)

        return parsed.isoformat()

    @staticmethod
    def parse_localized(kind: str, text: str) -> Any:
        """A date as the reader's language writes it: 26/09/2026."""
        names = {
            "date": ("DATE_INPUT_FORMATS",),
            "datetime": ("DATETIME_INPUT_FORMATS", "DATE_INPUT_FORMATS"),
            "time": ("TIME_INPUT_FORMATS",),
        }[kind]

        for setting in names:
            for pattern in formats.get_format(setting):
                try:
                    parsed = datetime.datetime.strptime(text, pattern)
                except (TypeError, ValueError):
                    continue

                return parsed.time() if kind == "time" else parsed

        return None

    def convert_number(self, column: ImportColumn, value: Any) -> Any:
        if isinstance(value, bool):
            value = int(value)

        if isinstance(value, (int, float, decimal.Decimal)):
            number: Any = value
        else:
            text = re.sub(r"[\s  ']", "", force_str(value))

            # 1.234,56 and 1,234.56 alike: the last separator is the
            # decimal one.
            if "," in text and "." in text:
                if text.rfind(",") > text.rfind("."):
                    text = text.replace(".", "").replace(",", ".")
                else:
                    text = text.replace(",", "")
            else:
                text = text.replace(",", ".")

            try:
                number = decimal.Decimal(text)
            except decimal.InvalidOperation:
                raise RowError(
                    gettext("%(value)s is not a number.")
                    % {"value": force_str(value)},
                    column.name,
                )

        if column.kind == "integer":
            if int(number) != number:
                raise RowError(
                    gettext("%(value)s is not a whole number.")
                    % {"value": force_str(value)},
                    column.name,
                )

            return int(number)

        if column.kind == "float":
            return float(number)

        return str(decimal.Decimal(str(number)))

    def related_queryset(self, column: ImportColumn) -> Any:
        """The records a cell may name: those the importer may see."""
        related = column.field.related_model
        resource = self.resource.site.get_resource(related)

        if resource is None:
            return related._default_manager.all()

        if not resource.has_view_permission(self.request):
            return related._default_manager.none()

        return resource.get_queryset(self.request)

    def find(self, column: ImportColumn, value: Any) -> Any:
        """The primary key of the one record a cell names."""
        if isinstance(value, float) and value.is_integer():
            value = int(value)

        text = force_str(value).strip()
        cache_key = (column.name, simplify(text))

        if cache_key in self._lookups:
            found = self._lookups[cache_key]
        else:
            queryset = self.related_queryset(column)
            lookup = column.lookup or "pk"

            if lookup == "pk":
                condition = {"pk": text}
            else:
                condition = {text_lookup(lookup, "iexact"): text}

            try:
                found = list(
                    queryset.filter(**condition).values_list("pk")[:2]
                )
            except (ValueError, ValidationError):
                found = []

            self._lookups[cache_key] = found

        label = force_str(column.field.related_model._meta.verbose_name)

        if not found:
            raise RowError(
                gettext("No %(model)s is called %(value)s.")
                % {"model": label, "value": text},
                column.name,
            )

        if len(found) > 1:
            raise RowError(
                gettext("Several %(model)s are called %(value)s.")
                % {
                    "model": force_str(
                        column.field.related_model._meta.verbose_name_plural
                    ),
                    "value": text,
                },
                column.name,
            )

        return found[0][0]

    # -- a row -----------------------------------------------------------

    def get_defaults(self) -> dict[str, Any]:
        defaults = self.declaration.defaults

        if callable(defaults):
            defaults = defaults(self.request)

        return dict(defaults or {})

    def find_existing(self, values: dict[str, Any]) -> Any:
        key = self.declaration.key

        if not key or values.get(key) in (None, ""):
            return None

        return (
            self.resource.get_queryset(self.request)
            .filter(**{key: values[key]})
            .first()
        )

    @staticmethod
    def is_unchanged(instance: Any, validated: dict[str, Any]) -> bool:
        for name, value in validated.items():
            field = model_field(type(instance), name)

            if field is None:
                return False

            if field.many_to_many:
                current = {obj.pk for obj in getattr(instance, name).all()}
                wanted = {getattr(obj, "pk", obj) for obj in value}

                if current != wanted:
                    return False

                continue

            current = getattr(instance, field.attname)
            wanted = (
                getattr(value, "pk", value) if field.is_relation else value
            )

            if current != wanted:
                return False

        return True

    def process(
        self,
        mapping: Sequence[str | None],
        cells: Sequence[Any],
        number: int,
    ) -> tuple[str, dict[str, Any]]:
        """Convert, validate and save one row: what was done to it."""
        values: dict[str, Any] = {}
        errors: list[RowError] = []

        for index, name in enumerate(mapping):
            if name is None:
                continue

            cell = cells[index] if index < len(cells) else None

            if is_blank(cell):
                continue

            try:
                values[name] = self.convert(self.column(name), cell)
            except RowError as error:
                errors.append(error)

        if errors:
            raise RowErrors(errors)

        try:
            values = self.resource.clean_import_row(
                self.request, values, number
            )
        except (ValidationError, ApiValidationError) as error:
            raise RowErrors([RowError(message) for message in messages(error)])

        instance = self.find_existing(values)
        declaration = self.declaration

        if instance is None:
            if not declaration.creates:
                raise RowErrors(
                    [
                        RowError(
                            gettext("No record has this %(key)s.")
                            % {"key": self.column(declaration.key).title},
                            declaration.key,
                        )
                    ]
                )

            if not self.resource.has_add_permission(self.request):
                raise RowErrors(
                    [RowError(gettext("You may not add records."))]
                )

            data = {**self.get_defaults(), **values}
            serializer = self.resource.get_form_serializer_class()(
                data=data, context={"request": self.request}
            )
            action = "create"
        else:
            if not declaration.updates:
                raise RowErrors(
                    [
                        RowError(
                            gettext("This record exists already."),
                            declaration.key,
                        )
                    ]
                )

            if not self.resource.has_change_permission(self.request, instance):
                raise RowErrors(
                    [RowError(gettext("You may not change this record."))]
                )

            serializer = self.resource.get_form_serializer_class()(
                instance,
                data=values,
                partial=True,
                context={"request": self.request},
            )
            action = "update"

        if not serializer.is_valid():
            raise RowErrors(
                [
                    RowError(message, field if field in values else None)
                    for field, message in flatten(serializer.errors, self)
                ]
            )

        if action == "update" and self.is_unchanged(
            instance, serializer.validated_data
        ):
            return "unchanged", values

        try:
            with transaction.atomic(using=self.database):
                self.resource.save_import_row(
                    self.request, serializer, instance
                )
        except IntegrityError:
            raise RowErrors(
                [RowError(gettext("The database refused this row."))]
            )

        return action, values

    @property
    def database(self) -> str:
        return router.db_for_write(self.resource.model)

    # -- the whole file --------------------------------------------------

    def run(
        self,
        upload: Any,
        raw_mapping: Any = None,
        commit: bool = False,
    ) -> dict[str, Any]:
        """Read, map, convert, validate and - asked to, and clean - write.

        Raises ``ImportRefused`` for a file that cannot be read at all;
        everything wrong with its rows comes back in the result.
        """
        from generic.history import acting_as
        from generic.sites.realtime import batch

        sheet = read_file(upload, self.max_rows)
        mapping = self.read_mapping(sheet.headers, raw_mapping)
        counts = {"create": 0, "update": 0, "unchanged": 0, "error": 0}
        errors: list[dict[str, Any]] = []
        preview: list[dict[str, Any]] = []
        preview_limit = generic_settings.IMPORT_PREVIEW_ROWS
        source = force_str(getattr(upload, "name", "") or "")

        with transaction.atomic(using=self.database):
            with acting_as(self.request.user, source=f"Import {source}"):
                with batch(self.resource):
                    for number, cells in sheet.rows:
                        try:
                            action, _values = self.process(
                                mapping, cells, number
                            )
                        except RowErrors as failure:
                            action = "error"

                            for error in failure.errors:
                                if len(errors) < MAX_ERRORS:
                                    errors.append(
                                        {
                                            "row": number,
                                            "column": error.column,
                                            "message": error.message,
                                        }
                                    )

                        counts[action] += 1

                        if len(preview) < preview_limit:
                            preview.append(
                                {
                                    "row": number,
                                    "action": action,
                                    "cells": [
                                        (
                                            display(cells[index])
                                            if index < len(cells)
                                            else ""
                                        )
                                        for index in range(len(mapping))
                                    ],
                                }
                            )

            committed = bool(commit and not counts["error"])

            # The preview, and a file with an error anywhere: nothing of
            # it stays.
            if not committed:
                transaction.set_rollback(True, using=self.database)

        return {
            "headers": sheet.headers,
            "mapping": mapping,
            "unmatched": [
                header
                for header, name in zip(sheet.headers, mapping)
                if name is None and header
            ],
            "missing": self.missing_columns(mapping),
            "rows": len(sheet.rows),
            "counts": counts,
            "errors": errors,
            "moreErrors": len(errors) >= MAX_ERRORS,
            "preview": preview,
            "committed": committed,
        }

    def missing_columns(self, mapping: Sequence[str | None]) -> list[str]:
        """Required columns no header goes to, when rows may be new."""
        if not self.declaration.creates:
            return []

        mapped = {name for name in mapping if name}
        defaults = self.get_defaults()

        return [
            column.name
            for column in self.get_columns()
            if column.required
            and column.name not in mapped
            and column.name not in defaults
        ]

    # -- the template ----------------------------------------------------

    def template(self) -> bytes:
        """An empty workbook with the headers, and the allowed values."""
        from openpyxl import Workbook

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = force_str(self.resource.get_label_plural())[:31]
        columns = self.get_columns()
        sheet.append([column.title for column in columns])

        listed = [column for column in columns if column.choices]

        if listed:
            values = workbook.create_sheet(gettext("Allowed values")[:31])
            values.append([column.title for column in listed])

            longest = max(len(column.choices) for column in listed)

            for index in range(longest):
                values.append(
                    [
                        (
                            list(column.choices.values())[index]
                            if index < len(column.choices)
                            else None
                        )
                        for column in listed
                    ]
                )

        stream = io.BytesIO()
        workbook.save(stream)

        return stream.getvalue()


class RowErrors(Exception):
    def __init__(self, errors: list[RowError]) -> None:
        super().__init__(errors)
        self.errors = errors


def display(value: Any) -> str:
    """A cell as the preview shows it."""
    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        value = int(value)

    if isinstance(value, datetime.datetime):
        return formats.localize(value)

    if isinstance(value, datetime.date):
        return formats.localize(value)

    return force_str(value)


def messages(error: Any) -> list[str]:
    detail = getattr(error, "detail", None)

    if detail is None:
        detail = getattr(error, "messages", [str(error)])

    return [message for _field, message in flatten(detail, None)]


def flatten(detail: Any, importer: Any) -> Iterable[tuple[str | None, str]]:
    """Serializer errors as ``(field, message)`` pairs, labelled."""
    if isinstance(detail, dict):
        for field, value in detail.items():
            name = None if field in ("non_field_errors", "__all__") else field

            for _inner, message in flatten(value, importer):
                if (
                    name
                    and importer is not None
                    and not any(
                        column.name == name
                        for column in importer.get_columns()
                    )
                ):
                    message = f"{field}: {message}"

                yield name, message
    elif isinstance(detail, (list, tuple)):
        for item in detail:
            yield from flatten(item, importer)
    else:
        yield None, force_str(detail)
